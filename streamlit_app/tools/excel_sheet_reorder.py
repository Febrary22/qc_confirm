"""엑셀 시트 순서를 바꾸는 도구."""
from __future__ import annotations

import io

import pandas as pd
import streamlit as st
from openpyxl import load_workbook

from common import download_result, friendly_errors, step_caption, tool_header

TOOL_META = {
    "name": "엑셀 시트 순서 바꾸기",
    "category": "Excel",
    "description": "엑셀 파일 안에 있는 시트들의 순서를 원하는 대로 바꿔드려요.",
    "icon": "🔀",
}


def render() -> None:
    tool_header(TOOL_META["name"], TOOL_META["description"])

    step_caption(1, "순서를 바꿀 엑셀 파일을 올려주세요")
    file = st.file_uploader("엑셀 파일 선택", type=["xlsx"])
    if not file:
        st.info("엑셀 파일을 1개 올려주세요. (.xlsx 형식만 가능해요)")
        return

    with friendly_errors("엑셀 시트 목록 확인"):
        file.seek(0)
        wb = load_workbook(file)
        sheet_names = wb.sheetnames

    step_caption(2, "새로운 순서 번호를 입력해주세요")
    order_df = pd.DataFrame(
        {
            "시트 이름": sheet_names,
            "새 순서": list(range(1, len(sheet_names) + 1)),
        }
    )
    edited = st.data_editor(
        order_df,
        hide_index=True,
        use_container_width=True,
        disabled=["시트 이름"],
        key="excel_reorder_editor",
    )

    step_caption(3, "실행 버튼을 눌러주세요")
    if st.button("시트 순서 바꾸기 실행", type="primary"):
        with friendly_errors("엑셀 시트 순서 바꾸기"):
            new_order = edited["새 순서"].tolist()
            if sorted(new_order) != list(range(1, len(sheet_names) + 1)):
                raise ValueError(
                    f"새 순서 번호는 1부터 {len(sheet_names)}까지 한 번씩만 사용해야 해요. "
                    "중복되거나 빠진 번호가 없는지 확인해 주세요."
                )

            order_map = list(zip(edited["시트 이름"].tolist(), new_order))
            order_map.sort(key=lambda pair: pair[1])
            ordered_names = [name for name, _ in order_map]

            file.seek(0)
            wb = load_workbook(file)
            wb._sheets = [wb[name] for name in ordered_names]  # noqa: SLF001

            buffer = io.BytesIO()
            wb.save(buffer)

            step_caption(4, "결과 파일을 받아주세요")
            download_result(
                "순서가 바뀐 엑셀 내려받기",
                buffer.getvalue(),
                "시트순서변경_결과.xlsx",
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
